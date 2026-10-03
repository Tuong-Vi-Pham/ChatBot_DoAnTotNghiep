import json
import os
from pathlib import Path
from docx import Document
from docx.shared import Inches

def load_markdown_tables(md_path: Path):
    """Extract markdown tables from the evaluation results markdown file.
    Returns a list of tables, each as a list of rows (list of strings)."""
    tables = []
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
    current = []
    in_table = False
    for line in lines:
        if line.strip().startswith("|"):
            in_table = True
            current.append(line.strip())
        else:
            if in_table:
                tables.append(current)
                current = []
                in_table = False
    if current:
        tables.append(current)
    return tables

def add_markdown_table_to_doc(doc: Document, table_lines):
    # First line is header, second line separators, rest data
    header = [h.strip() for h in table_lines[0].split("|") if h]
    rows = []
    for line in table_lines[2:]:
        cols = [c.strip() for c in line.split("|") if c]
        rows.append(cols)
    table = doc.add_table(rows=1, cols=len(header))
    hdr_cells = table.rows[0].cells
    for idx, text in enumerate(header):
        hdr_cells[idx].text = text
    for row in rows:
        row_cells = table.add_row().cells
        for idx, text in enumerate(row):
            row_cells[idx].text = text
    doc.add_paragraph()  # spacing

def main():
    repo_root = Path(__file__).resolve().parents[1]
    data_path = repo_root / "evaluation" / "report_data.json"
    md_path = repo_root / "evaluation" / "evaluation_results.md"
    out_path = repo_root / "evaluation" / "SmartLogi_Experimental_Evaluation_Report.docx"

    # Load data
    with open(data_path, "r", encoding="utf-8") as f:
        test_entries = json.load(f)

    doc = Document()
    doc.add_heading("SmartLogi Experimental Evaluation Report", 0)

    # 1. Introduction
    doc.add_heading("1. Introduction & Objectives", level=1)
    doc.add_paragraph(
        "This report presents the experimental evaluation of the SmartLogi chatbot system. "
        "A total of 60 test cases were executed, and the results are summarized below."
    )

    # 2. Experimental Setup (placeholder – user can edit later)
    doc.add_heading("2. Experimental Setup", level=1)
    doc.add_paragraph(
        "*Environment details (Python version, OS, hardware) should be inserted here."
    )

    # 3. Test Case Table
    doc.add_heading("3. Test Cases", level=1)
    # Create table with selected columns
    columns = ["Test_ID", "Query", "Category", "Expected_Answer", "System_Response", "Judge_1", "Judge_2", "Judge_3", "Final_Label", "Agreement"]
    table = doc.add_table(rows=1, cols=len(columns))
    hdr_cells = table.rows[0].cells
    for i, col in enumerate(columns):
        hdr_cells[i].text = col
    for entry in test_entries:
        row_cells = table.add_row().cells
        for i, col in enumerate(columns):
            val = entry.get(col) or ""
            # Truncate very long responses for readability
            if isinstance(val, str) and len(val) > 200:
                val = val[:197] + "..."
            row_cells[i].text = str(val)
    doc.add_paragraph()

    # 4. Metric Tables
    doc.add_heading("4. Evaluation Metrics", level=1)
    tables = load_markdown_tables(md_path)
    for idx, tbl in enumerate(tables, start=1):
        doc.add_heading(f"Metric Table {idx}", level=2)
        add_markdown_table_to_doc(doc, tbl)

    # 5. Analysis & Discussion (placeholder)
    doc.add_heading("5. Analysis & Discussion", level=1)
    doc.add_paragraph("*Insert analysis of results, observations, and discussion here.*")

    # 6. Conclusion
    doc.add_heading("6. Conclusion & Future Work", level=1)
    doc.add_paragraph("*Summarize key findings and outline potential future improvements.*")

    # 7. Appendix
    doc.add_heading("7. Appendix", level=1)
    doc.add_paragraph("Full test case data is available in `evaluation/report_data.json`.")

    doc.save(out_path)
    print(f"Report generated at {out_path}")

if __name__ == "__main__":
    main()
