import torch  # Critical: Import torch first to avoid shm.dll conflicts on Windows CPU
import os
import tempfile
from typing import List, Optional, Any
from src.loaders.document import Document

class DocumentLoader:
    """
    Loader for team documents including PDF, DOCX, TXT, and images (PNG, JPG).
    Utilizes PaddleOCR for image OCR and scanned PDF pages.
    """
    def __init__(self, show_ocr_log: bool = False):
        self._ocr: Optional[Any] = None
        self.show_ocr_log = show_ocr_log

    def _get_ocr_engine(self) -> Any:
        """
        Lazy initializer for PaddleOCR to avoid loading it in memory unless needed.
        """
        if self._ocr is None:
            try:
                from paddleocr import PaddleOCR
                self._ocr = PaddleOCR(lang='en')
            except ImportError:
                print("[DocumentLoader Warning] paddleocr module not found; OCR fallback disabled.")
                return None
        return self._ocr

    def load_txt(self, file_path: str) -> str:
        """Extract text from TXT files."""
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()

    def load_docx(self, file_path: str) -> str:
        """Extract text from DOCX files, including paragraphs and tables."""
        try:
            import docx
        except ImportError:
            print("[DocumentLoader Warning] python-docx module not found.")
            return ""

        doc = docx.Document(file_path)
        content_parts = []
        
        # Extract paragraph texts
        for para in doc.paragraphs:
            if para.text.strip():
                content_parts.append(para.text.strip())
                
        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    content_parts.append(" | ".join(cells))
                    
        return "\n".join(content_parts)

    def load_image(self, file_path: str) -> str:
        """Extract text from image files (JPG, PNG) using PaddleOCR."""
        ocr_engine = self._get_ocr_engine()
        result = ocr_engine.ocr(file_path, cls=True)
        
        text_lines = []
        if result:
            for res in result:
                if res:
                    for line in res:
                        if line and len(line) > 1 and len(line[1]) > 0:
                            text_lines.append(line[1][0])
                            
        return "\n".join(text_lines)

    def load_pdf(self, file_path: str) -> str:
        """
        Extract text from PDF files.
        First attempts native text extraction. If text is insufficient or empty
        (indicating a scanned PDF), renders each page to an image and runs PaddleOCR.
        """
        try:
            import fitz
        except ImportError:
            print("[DocumentLoader Warning] PyMuPDF (fitz) module not found.")
            return ""

        doc = fitz.open(file_path)
        native_text_parts = []
        
        # 1. Native text extraction attempt
        for page in doc:
            page_text = page.get_text()
            if page_text.strip():
                native_text_parts.append(page_text.strip())
                
        full_native_text = "\n".join(native_text_parts).strip()
        
        # Check if the extracted text length is sufficient
        # If it is less than 50 characters, we treat it as a scanned PDF
        if len(full_native_text) > 50:
            return full_native_text
            
        # 2. Fallback to OCR for scanned PDF
        ocr_text_parts = []
        ocr_engine = self._get_ocr_engine()
        
        with tempfile.TemporaryDirectory() as temp_dir:
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                # Render page at 150 DPI for accurate text extraction
                pix = page.get_pixmap(dpi=150)
                temp_img_path = os.path.join(temp_dir, f"page_{page_num}.png")
                pix.save(temp_img_path)
                
                # Run PaddleOCR on the page image
                result = ocr_engine.ocr(temp_img_path, cls=True)
                page_lines = []
                if result:
                    for res in result:
                        if res:
                            for line in res:
                                if line and len(line) > 1 and len(line[1]) > 0:
                                    page_lines.append(line[1][0])
                
                if page_lines:
                    ocr_text_parts.append("\n".join(page_lines))
                    
        return "\n\n".join(ocr_text_parts)

    def load_file(self, file_path: str, base_dir: Optional[str] = None) -> Optional[Document]:
        """
        Detect file extension and load text, returning a structured Document object.
        """
        filename = os.path.basename(file_path)
        # Skip temporary Office lock files (~$) and hidden files (.)
        if filename.startswith('~$') or filename.startswith('.'):
            return None

        ext = os.path.splitext(file_path)[1].lower()
        if ext not in {'.pdf', '.docx', '.txt', '.png', '.jpg', '.jpeg'}:
            return None
            
        try:
            if ext == '.txt':
                text = self.load_txt(file_path)
            elif ext == '.docx':
                text = self.load_docx(file_path)
            elif ext in {'.png', '.jpg', '.jpeg'}:
                text = self.load_image(file_path)
            elif ext == '.pdf':
                text = self.load_pdf(file_path)
            else:
                return None
                
            text = text.strip()
            if not text:
                return None
                
            # Derive category and relative_path
            parent_dir = os.path.basename(os.path.dirname(file_path))
            if base_dir:
                try:
                    rel_path = os.path.relpath(file_path, base_dir)
                except ValueError:
                    rel_path = file_path
            else:
                rel_path = file_path

            # Normalize path slashes to forward slashes for consistent hashing/IDs across OS
            rel_path_norm = rel_path.replace('\\', '/')
            
            metadata = {
                "source_type": "document",
                "category": parent_dir,
                "source": filename,
                "file_name": filename,
                "full_path": os.path.abspath(file_path),
                "relative_path": rel_path_norm,
                "document_type": ext.lstrip('.')
            }
            
            return Document(page_content=text, metadata=metadata)
        except Exception as e:
            print(f"Error loading file {file_path}: {e}")
            return None

    def load_directory(self, dir_path: str, base_dir: Optional[str] = None) -> List[Document]:
        """
        Traverse directory recursively to load all supported documents.
        """
        documents = []
        effective_base = base_dir if base_dir is not None else dir_path
        for root, _, files in os.walk(dir_path):
            for file in sorted(files):
                if file.startswith('~$') or file.startswith('.'):
                    continue
                file_path = os.path.join(root, file)
                doc = self.load_file(file_path, base_dir=effective_base)
                if doc:
                    documents.append(doc)
        return documents
